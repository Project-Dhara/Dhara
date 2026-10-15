"""Data ingestion + spreadsheet upload pipeline (Excel / CSV). PDF tests optional."""
from __future__ import annotations

import io
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

from app.services.ingestion import csv as ingest_csv
from app.services.ingestion import excel as ingest_excel
from app.services.ingestion.service import run_upload_pipeline
from app.services.ingestion.spreadsheet_pipeline import (
    ingest_spreadsheet_path,
    run_validation_from_pre_input,
)
from app.services.ingestion.utils.source_type import resolve_source_type

_REPO_ROOT = Path(__file__).resolve().parents[2]
_TESTING_FILES = _REPO_ROOT / "Testing files"


def _sandbox_excel_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data points"
    ws.append(["Serial Number", "Table Name", "Variables (Summary)", "Unit"])
    ws.append([1, "Scheme Summary", "Scheme name, beneficiaries", "Scheme"])
    ws.append([2, "KPI Monitoring", "Applications, approvals", "District"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _multi_table_sheet_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Combined Tables"
    ws.append([None] * 6)
    ws.append(["D-14"] * 6)
    ws.append(["INFANT DEATHS BY AGE (URBAN)"] * 6)
    ws.append(["SL. NO.", "AGE", "MALE", "FEMALE", "OTHER", "TOTAL"])
    ws.append([1, "< 7 DAYS", 10, 8, 0, 18])
    ws.append([2, "7-28 DAYS", 5, 4, 0, 9])
    ws.append([None] * 6)
    ws.append([None] * 6)
    ws.append(["D-18"] * 6)
    ws.append(["PREGNANCY RELATED DEATHS (URBAN)"] * 6)
    ws.append(["SL. NO.", "OCCUPATION", "<15", "15-19", "20-24", "25-29"])
    ws.append([1, "PROFESSIONAL", 0, 0, 2, 1])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _titled_sheet_bytes() -> bytes:
    """Merged/fill-expanded title banner — must not become the sheet name."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MySheetName"
    ws.append(["Scheme Summary (Basic Statistics)"] * 4)
    ws.append(["Name", "Beneficiaries", "Budget", "Year"])
    ws.append(["S1", 10, 100, 2020])
    ws.append(["S2", 20, 200, 2021])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sample_csv_bytes() -> bytes:
    return (
        "Year,Population,Birth Rate\n"
        "2020,1000,15.2\n"
        "2021,1020,14.8\n"
    ).encode("utf-8")


def _assert_pre_validation_shape(pre: dict) -> None:
    assert "pages_grouped" in pre and "unit_context" in pre
    grouped = pre["pages_grouped"]
    ctx = pre["unit_context"]
    assert grouped
    for unit_id, cands in grouped.items():
        assert unit_id in ctx
        assert ctx[unit_id]
        assert len(cands) >= 1
        assert cands[0].get("df") is not None


def _table_count(result: dict) -> int:
    return sum(len(page.get("tables") or []) for page in (result or {}).values())


def _assert_validated_table_shape(result: dict) -> None:
    assert _table_count(result) >= 1
    for page in result.values():
        for t in page.get("tables") or []:
            cols = t.get("columns") or []
            ncol = len(cols)
            assert ncol >= 1
            for row in t.get("rows") or []:
                assert len(row) == ncol


def test_resolve_source_type_from_extension():
    assert resolve_source_type("report.pdf", None) == "pdf"
    assert resolve_source_type("data.xlsx", "excel") == "excel"
    assert resolve_source_type("export.csv", "csv") == "csv"


def test_excel_pre_validation_input():
    pre = ingest_excel.ingest_bytes(_sandbox_excel_bytes(), "sandbox.xlsx")
    _assert_pre_validation_shape(pre)
    assert pre.get("source_type") == "excel"
    assert pre["pages_grouped"][1][0]["method"] == "excel_grid"
    assert pre["pages_grouped"][1][0]["sheet"] == "Data points"


def test_excel_multi_table_pre_validation_input():
    pre = ingest_excel.ingest_bytes(_multi_table_sheet_bytes(), "multi.xlsx")
    _assert_pre_validation_shape(pre)
    assert len(pre["pages_grouped"]) >= 2


def test_excel_preserves_in_grid_title_not_sheet_name():
    """Merge-filled title banners are stripped for headers but kept as caption."""
    pre = ingest_excel.ingest_bytes(_titled_sheet_bytes(), "demo.xlsx")
    _assert_pre_validation_shape(pre)
    cand = pre["pages_grouped"][1][0]
    assert cand["sheet"] == "MySheetName"
    caption = (getattr(cand["df"], "attrs", {}) or {}).get("dhara_table_caption") or ""
    assert caption == "Scheme Summary (Basic Statistics)"
    assert "MySheetName" not in caption
    # Context lists the real title before sheet metadata.
    ctx = pre["unit_context"][1]
    assert "Scheme Summary (Basic Statistics)" in ctx.splitlines()[0]
    assert ctx.strip().endswith("Sheet: MySheetName")


def test_excel_upload_title_is_not_sheet_name(tmp_path):
    path = tmp_path / "titled.xlsx"
    path.write_bytes(_titled_sheet_bytes())
    result = run_upload_pipeline(path, "excel", api_key=None)
    _assert_validated_table_shape(result)
    titles = [t.get("title") for page in result.values() for t in page["tables"]]
    assert titles
    assert titles[0] == "Scheme Summary (Basic Statistics)"
    assert all(t != "MySheetName" for t in titles)
    assert all(not str(t or "").startswith("Sheet:") for t in titles)
    assert all(not str(t or "").startswith("Source file:") for t in titles)


def test_excel_upload_short_code_banner_keeps_descriptive_title(tmp_path):
    """Fill-expanded ``D-14`` + title banners must not yield the sheet name."""
    path = tmp_path / "multi.xlsx"
    path.write_bytes(_multi_table_sheet_bytes())
    result = run_upload_pipeline(path, "excel", api_key=None)
    _assert_validated_table_shape(result)
    titles = [t.get("title") for page in result.values() for t in page["tables"]]
    assert titles
    assert any("INFANT DEATHS" in str(t or "").upper() for t in titles)
    assert all(t != "Combined Tables" for t in titles)
    assert all(str(t or "").upper() != "D-14" for t in titles)


def test_csv_pre_validation_input():
    pre = ingest_csv.ingest_bytes(_sample_csv_bytes(), "sample.csv")
    _assert_pre_validation_shape(pre)
    assert pre.get("source_type") == "csv"
    assert pre["pages_grouped"][1][0]["method"] == "csv_grid"


def test_spreadsheet_validation_from_pre_input(tmp_path):
    path = tmp_path / "sample.csv"
    path.write_bytes(_sample_csv_bytes())
    pre = ingest_spreadsheet_path(path, "csv")
    result = run_validation_from_pre_input(pre, "csv", api_key=None)
    _assert_validated_table_shape(result)
    assert result[1]["tables"][0].get("source_type") == "csv"


def test_run_upload_pipeline_excel_no_llm_key(tmp_path):
    path = tmp_path / "sandbox.xlsx"
    path.write_bytes(_sandbox_excel_bytes())
    result = run_upload_pipeline(path, "excel", api_key=None)
    _assert_validated_table_shape(result)
    for page in result.values():
        for t in page["tables"]:
            assert t.get("source_type") == "excel"
            assert str(t.get("sheet") or "").strip(), "sheet must survive validation for Preview subtitle"


def test_stamp_bbox_copies_sheet_from_excel_candidate():
    """LLM-validated tables must keep sheet so ambiguous-column cards show it."""
    from app.services.extraction.pdf.confidence_validation.alignment import (
        _stamp_bbox_from_candidate,
    )
    table: dict = {"title": "Demo"}
    _stamp_bbox_from_candidate(table, {"sheet": "Vital Stats", "source_file": "a.xlsx"})
    assert table["sheet"] == "Vital Stats"
    assert table["source_file"] == "a.xlsx"


def test_run_upload_pipeline_csv_no_llm_key(tmp_path):
    path = tmp_path / "sample.csv"
    path.write_bytes(_sample_csv_bytes())
    result = run_upload_pipeline(path, "csv", api_key=None)
    _assert_validated_table_shape(result)


@pytest.mark.skipif(
    not (_TESTING_FILES / "SDA_INDIA_222-224.pdf").is_file(),
    reason="Testing files/SDA_INDIA_222-224.pdf not present",
)
def test_run_upload_pipeline_pdf_unchanged(tmp_path):
    """PDF branch still matches direct run_pipeline (not PreValidationInput yet)."""
    pdf_path = _TESTING_FILES / "SDA_INDIA_222-224.pdf"
    from app.services.extraction.pdf import confidence_validation as pe

    direct = pe.run_pipeline(pdf_path, api_key=None)
    via_upload = run_upload_pipeline(pdf_path, "pdf", api_key=None)
    assert set(direct.keys()) == set(via_upload.keys())
    assert _table_count(direct) == _table_count(via_upload)


if __name__ == "__main__":
    test_resolve_source_type_from_extension()
    test_excel_pre_validation_input()
    test_excel_multi_table_pre_validation_input()
    test_csv_pre_validation_input()
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        xlsx = root / "sandbox.xlsx"
        xlsx.write_bytes(_sandbox_excel_bytes())
        csv_path = root / "sample.csv"
        csv_path.write_bytes(_sample_csv_bytes())
        pre = ingest_spreadsheet_path(xlsx, "excel")
        run_validation_from_pre_input(pre, "excel", api_key=None)
        _assert_validated_table_shape(run_upload_pipeline(xlsx, "excel", api_key=None))
        _assert_validated_table_shape(run_upload_pipeline(csv_path, "csv", api_key=None))
    print("test_data_ingestion_pipeline: OK")
