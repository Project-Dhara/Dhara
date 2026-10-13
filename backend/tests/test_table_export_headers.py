"""Downloaded .xlsx headers must match Preview multi-level layout."""
from __future__ import annotations

from io import BytesIO

import openpyxl

from app.services.ingestion.wizard.table_export import (
    _header_paths_for_table,
    table_to_excel_bytes,
)


def _load(table: dict):
    return openpyxl.load_workbook(BytesIO(table_to_excel_bytes(table)))


def test_pdf_table_exports_merged_group_headers():
    table = {
        "title": "Live births by area",
        "columns": [
            {"name": "State", "header_path": ["State"]},
            {"name": "Male", "header_group": "URBAN", "header_path": ["URBAN", "Male"]},
            {"name": "Female", "header_group": "URBAN", "header_path": ["URBAN", "Female"]},
            {"name": "Male (2)", "header_group": "RURAL", "header_path": ["RURAL", "Male"]},
            {"name": "Female (2)", "header_group": "RURAL", "header_path": ["RURAL", "Female"]},
        ],
        "rows": [
            ["01", "10", "20", "30", "40"],
            ["02", "11", "21", "31", "41"],
        ],
    }
    wb = _load(table)
    ws = wb.active

    # Preview shows URBAN spanning Male|Female (cols B–C), RURAL spanning D–E
    assert ws.cell(1, 1).value == "State"
    assert ws.cell(1, 2).value == "URBAN"
    assert ws.cell(1, 4).value == "RURAL"
    assert ws.cell(2, 2).value == "Male"
    assert ws.cell(2, 3).value == "Female"
    assert ws.cell(2, 4).value == "Male"
    assert ws.cell(2, 5).value == "Female"
    # Leaf labels come from header_path, not uniquified name
    assert ws.cell(2, 4).value == "Male"
    assert ws.cell(3, 1).value == "01"
    assert ws.cell(3, 2).value == "10"

    ranges = {str(r) for r in ws.merged_cells.ranges}
    assert "B1:C1" in ranges  # URBAN
    assert "D1:E1" in ranges  # RURAL
    assert "A1:A2" in ranges  # State stub spans both header rows


def test_excel_raw_header_rows_drive_export():
    table = {
        "title": "Area table",
        "columns": ["State", "URBAN / Male", "URBAN / Female", "RURAL / Male", "RURAL / Female"],
        "raw_header_rows": [
            ["State", "URBAN", "URBAN", "RURAL", "RURAL"],
            ["State", "Male", "Female", "Male", "Female"],
        ],
        "rows": [
            {"State": "01", "URBAN / Male": 10, "URBAN / Female": 20, "RURAL / Male": 30, "RURAL / Female": 40},
        ],
    }
    paths = _header_paths_for_table(table, table["columns"])
    assert paths[1] == ["URBAN", "Male"]
    assert paths[3] == ["RURAL", "Male"]

    wb = _load(table)
    ws = wb.active
    assert ws.cell(1, 2).value == "URBAN"
    assert ws.cell(2, 2).value == "Male"
    ranges = {str(r) for r in ws.merged_cells.ranges}
    assert "B1:C1" in ranges


def test_flat_slash_names_become_multi_level():
    """Catalogue rebuilds often only have flattened 'Group / Leaf' names."""
    table = {
        "title": "From DB",
        "columns": ["State", "URBAN / Male", "URBAN / Female"],
        "rows": [
            {"State": "01", "URBAN / Male": 1, "URBAN / Female": 2},
        ],
    }
    wb = _load(table)
    ws = wb.active
    assert ws.cell(1, 2).value == "URBAN"
    assert ws.cell(2, 2).value == "Male"
    assert "B1:C1" in {str(r) for r in ws.merged_cells.ranges}


def test_single_level_headers_unchanged():
    table = {
        "title": "Simple",
        "columns": ["A", "B"],
        "rows": [{"A": 1, "B": 2}],
    }
    wb = _load(table)
    ws = wb.active
    assert ws.cell(1, 1).value == "A"
    assert ws.cell(1, 2).value == "B"
    assert ws.cell(2, 1).value == 1
    assert list(ws.merged_cells.ranges) == []
