"""Builds a single-sheet .xlsx from an extracted table so downloads match
Preview: multi-level column headers (header_path / raw_header_rows) with
merged group bands, then one data row per record.

Supports both Excel-pipeline tables (columns: list[str], rows: list[dict],
optional raw_header_rows) and PDF-pipeline tables (columns: list[{name,
header_path, …}], rows: list[list]).
"""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any, Optional

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


_BANNER_ROOT_RE = re.compile(
    r"^(?:table|statement|annex(?:ure)?)\s*[:.\-–—]?\s*\S",
    re.IGNORECASE,
)


def _column_names(table: dict) -> list[str]:
    columns = table.get("columns") or []
    names: list[str] = []
    for i, col in enumerate(columns):
        if isinstance(col, str):
            names.append(col or f"Col_{i + 1}")
        elif isinstance(col, dict):
            names.append(str(col.get("name") or f"Col_{i + 1}"))
        else:
            names.append(f"Col_{i + 1}")
    return names


def _iter_row_values(row: Any, col_names: list[str]) -> list[Any]:
    if isinstance(row, dict):
        return [row.get(c) for c in col_names]
    if isinstance(row, (list, tuple)):
        vals = list(row[: len(col_names)])
        while len(vals) < len(col_names):
            vals.append(None)
        return vals
    return [None] * len(col_names)


def _safe_sheet_title(name: Any) -> str:
    """Excel forbids \\ / ? * [ ] : in worksheet titles (max 31 chars)."""
    cleaned = re.sub(r'[:\\/?*\[\]]+', ' ', str(name or 'Table'))
    cleaned = re.sub(r'\s+', ' ', cleaned).strip(" '") or 'Table'
    return cleaned[:31] or 'Table'


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    return s


def _forward_fill_row(row: list[Any], n_cols: int) -> list[str]:
    filled: list[str] = []
    last = ""
    for i in range(n_cols):
        raw = _cell_str(row[i]) if i < len(row) else ""
        if raw:
            last = raw
            filled.append(raw)
        else:
            filled.append(last)
    return filled


def _collapse_consecutive(parts: list[str]) -> list[str]:
    out: list[str] = []
    for p in parts:
        s = (p or "").strip()
        if not s:
            continue
        if out and out[-1].lower() == s.lower():
            continue
        out.append(s)
    return out


def _filter_banner_roots(path: list[str]) -> list[str]:
    if not path:
        return path
    filtered = [
        p for i, p in enumerate(path)
        if i == len(path) - 1 or not _BANNER_ROOT_RE.match(p)
    ]
    return filtered or path[-1:]


def _column_header_path(col: Any) -> list[str]:
    """Mirror frontend columnHeaderPath — display path, not uniquified name."""
    if isinstance(col, str):
        name = col.strip()
        parts = [p.strip() for p in re.split(r"\s*/\s*", name) if p.strip()]
        return parts if parts else ([name] if name else [""])

    if not isinstance(col, dict):
        return [""]

    raw = col.get("header_path")
    if isinstance(raw, (list, tuple)) and len(raw) > 0:
        path = [_cell_str(p) for p in raw if _cell_str(p)]
        if path:
            return _filter_banner_roots(path)

    name = _cell_str(col.get("name"))
    group = _cell_str(col.get("header_group"))
    if group and group.lower() != name.lower() and not _BANNER_ROOT_RE.match(group):
        return [group, name] if name else [group]
    if name:
        parts = [p.strip() for p in re.split(r"\s*/\s*", name) if p.strip()]
        return parts if len(parts) >= 2 else [name]
    return [""]


def _paths_from_raw_header_rows(
    raw_header_rows: list,
    col_names: list[str],
) -> Optional[list[list[str]]]:
    n_cols = len(col_names)
    if n_cols <= 0:
        return None
    matrix = [
        r for r in (raw_header_rows or [])
        if isinstance(r, (list, tuple)) and any(_cell_str(v) for v in r)
    ]
    if len(matrix) < 2:
        return None
    filled = [_forward_fill_row(list(r), n_cols) for r in matrix]
    paths: list[list[str]] = []
    for i, name in enumerate(col_names):
        path = _collapse_consecutive([row[i] for row in filled])
        paths.append(_filter_banner_roots(path) if path else [name])
    return paths


def _header_paths_for_table(table: dict, col_names: list[str]) -> list[list[str]]:
    """Build per-column header paths matching Preview rendering."""
    if not col_names:
        return []

    from_raw = _paths_from_raw_header_rows(table.get("raw_header_rows") or [], col_names)
    if from_raw is not None:
        return from_raw

    columns = table.get("columns") or []
    if columns and all(isinstance(c, dict) for c in columns):
        paths = [_column_header_path(c) for c in columns]
        # Pad / trim to col_names length
        while len(paths) < len(col_names):
            paths.append([col_names[len(paths)]])
        return paths[: len(col_names)]

    return [_column_header_path(name) for name in col_names]


def _place_header_path(path: list[str], depth: int) -> list[Optional[dict]]:
    """Place a column path into a depth-tall grid (frontend placeHeaderPath)."""
    cells: list[Optional[dict]] = [None] * depth
    if not path:
        return cells
    extra = max(0, depth - len(path))
    row = 0
    for i, label in enumerate(path):
        row_span = (1 + extra) if i == 0 else 1
        prefix = "\0".join(path[: i + 1])
        cells[row] = {"label": label, "row_span": row_span, "prefix": prefix}
        for k in range(1, row_span):
            if row + k < depth:
                cells[row + k] = {"covered": True}
        row += row_span
    return cells


def _build_header_segments(paths: list[list[str]]) -> tuple[int, list[list[dict]]]:
    """Build N header rows of visible segments with colspan/rowspan."""
    if not paths:
        return 1, []
    depth = max(1, max((len(p) for p in paths), default=1))
    grid = [_place_header_path(p, depth) for p in paths]
    rows: list[list[dict]] = []
    n_cols = len(paths)
    for r in range(depth):
        segs: list[dict] = []
        c = 0
        while c < n_cols:
            cell = grid[c][r]
            if not cell or cell.get("covered"):
                c += 1
                continue
            col_span = 1
            while (
                c + col_span < n_cols
                and grid[c + col_span][r]
                and not grid[c + col_span][r].get("covered")
                and grid[c + col_span][r].get("prefix") == cell.get("prefix")
                and grid[c + col_span][r].get("row_span") == cell.get("row_span")
            ):
                col_span += 1
            segs.append({
                "label": cell["label"],
                "col_span": col_span,
                "row_span": int(cell["row_span"]),
                "start": c,
            })
            c += col_span
        rows.append(segs)
    return depth, rows


def _write_multi_level_headers(ws, paths: list[list[str]]) -> int:
    """Write Preview-matching header rows with merges. Returns header row count."""
    depth, rows = _build_header_segments(paths)
    if not paths:
        return 0

    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for r, segs in enumerate(rows):
        excel_row = r + 1
        for seg in segs:
            start_col = seg["start"] + 1
            end_col = seg["start"] + seg["col_span"]
            end_row = excel_row + seg["row_span"] - 1
            cell = ws.cell(row=excel_row, column=start_col, value=seg["label"])
            cell.font = header_font
            is_group = seg["col_span"] > 1 or (
                depth > 1 and excel_row + seg["row_span"] - 1 < depth
            )
            cell.alignment = center if is_group or seg["col_span"] > 1 else left
            if end_col > start_col or end_row > excel_row:
                ws.merge_cells(
                    start_row=excel_row,
                    start_column=start_col,
                    end_row=end_row,
                    end_column=end_col,
                )
    return depth


def table_to_excel_bytes(table: dict) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    sheet_name = table.get("sheet") or table.get("title") or table.get("table_id") or "Table"
    ws.title = _safe_sheet_title(sheet_name)

    col_names = _column_names(table)
    rows = table.get("rows") or []
    paths = _header_paths_for_table(table, col_names)

    header_depth = 0
    if paths:
        header_depth = _write_multi_level_headers(ws, paths)
    elif col_names:
        ws.append(col_names)
        header_depth = 1

    for row in rows:
        ws.append(_iter_row_values(row, col_names))

    # Modest column widths so multi-line headers stay readable
    for i in range(1, len(col_names) + 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = min(28, max(10, len(str(paths[i - 1][-1] if paths else col_names[i - 1])) + 2))

    if header_depth > 1:
        for r in range(1, header_depth + 1):
            ws.row_dimensions[r].height = 18

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def safe_download_stem(table: dict, fallback: str = "table") -> str:
    """Filesystem-safe basename (no extension) for a table download."""
    raw = (
        table.get("title")
        or table.get("table_id")
        or table.get("id")
        or fallback
    )
    stem = re.sub(r"[^\w.\-]+", "_", str(raw), flags=re.UNICODE).strip("._")
    return (stem or fallback)[:80]
